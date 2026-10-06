from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from datetime import timedelta
import pytest
from fastapi import HTTPException
from sqlalchemy import select,func
from aho.db import Session,now
from aho.models import *
from aho.schemas import *
from aho.services import *
from aho.queries import analytics,filtered
from conftest import login

def make(world,key=None,category=None):
    with Session.begin() as db:
        t=create_ticket(db,db.get(User,world['employee']),TicketCreate(category_id=category or world['category'],description='Не работает компьютер',building='Главный офис',floor='3',room='305'),key)
        return t.id

def act(world,identifier,who,action,**data):
    with Session.begin() as db: return action_ticket(db,db.get(User,world[who]),identifier,action,Action(**data)).status

def test_lifecycle_and_privacy(world):
    tid=make(world)
    assert act(world,tid,'akmal','accept')=='ACCEPTED'
    assert act(world,tid,'akmal','start')=='IN_PROGRESS'
    assert act(world,tid,'akmal','request-info',comment='Когда возникла проблема?')=='WAITING_REQUESTER'
    with Session.begin() as db:
        add_comment(db,db.get(User,world['employee']),tid,Comment(text='Сегодня утром'))
        add_comment(db,db.get(User,world['akmal']),tid,Comment(text='Внутренняя заметка',visibility='INTERNAL'))
    assert act(world,tid,'akmal','resume')=='IN_PROGRESS'
    assert act(world,tid,'akmal','complete',comment='Заменен блок питания')=='COMPLETED'
    assert act(world,tid,'employee','reopen',comment='Компьютер выключается')=='REOPENED'
    act(world,tid,'akmal','start');act(world,tid,'akmal','complete',comment='Проверили под нагрузкой')
    assert act(world,tid,'employee','confirm')=='CLOSED'
    act(world,tid,'employee','rate',score=5)
    with Session() as db:
        assert db.get(Rating,tid).score==5
        assert db.scalar(select(func.count()).select_from(StatusHistory).where(StatusHistory.ticket_id==tid))==10
        assert db.scalar(select(func.count()).select_from(Audit).where(Audit.entity_id==tid))>=12
        assert len(db.scalars(select(Notification).where(Notification.ticket_id==tid)).all())>8

@pytest.mark.parametrize('who,action,code',[('employee','accept',403),('employee2','accept',404),('akmal','complete',403),('manager','complete',409)])
def test_forbidden_actions(world,who,action,code):
    tid=make(world)
    with pytest.raises(HTTPException) as e: act(world,tid,who,action,comment='reason')
    assert e.value.status_code==code

def test_first_accept_wins_concurrently(world):
    tid=make(world); barrier=Barrier(2)
    def accept(who):
        with Session.begin() as db:
            u=db.get(User,world[who]);barrier.wait()
            try: action_ticket(db,u,tid,'accept',Action());return 200
            except HTTPException as e:return e.status_code
    with ThreadPoolExecutor(2) as pool: results=list(pool.map(accept,['akmal','javohir']))
    assert sorted(results)==[200,409]
    with Session() as db:
        winner=db.get(Ticket,tid).assigned_to
        assert winner in [world['akmal'],world['javohir']]
        assert db.scalar(select(func.count()).select_from(Assignment).where(Assignment.ticket_id==tid))==1
    loser='javohir' if winner==world['akmal'] else 'akmal'
    with pytest.raises(HTTPException):act(world,tid,loser,'accept')
    with Session() as db: assert db.get(Ticket,tid).assigned_to==winner

def test_idempotency_and_unique_numbers(world):
    assert make(world,'same')==make(world,'same')
    with ThreadPoolExecutor(4) as pool: ids=list(pool.map(lambda _:make(world),range(8)))
    with Session() as db:
        nums=list(db.scalars(select(Ticket.ticket_number).where(Ticket.id.in_(ids))))
        assert len(set(nums))==8
        assert all(n.startswith(f'AHO-{now().year}-') and len(n.split('-')[2])==6 for n in nums)

def test_cancellation_rules(world):
    tid=make(world);act(world,tid,'akmal','accept');act(world,tid,'akmal','start')
    assert act(world,tid,'employee','cancel',comment='Больше не требуется')=='IN_PROGRESS'
    with Session() as db:assert db.get(Ticket,tid).cancellation_requested
    assert act(world,tid,'akmal','approve-cancel')=='CANCELLED'
    tid=make(world);assert act(world,tid,'employee','cancel',comment='Ошибка')=='CANCELLED'

def test_sla_and_priority(world):
    tid=make(world)
    with Session.begin() as db:
        t=db.get(Ticket,tid);assert t.response_deadline-t.created_at==timedelta(hours=2)
        assert sla_state(t,t.created_at+timedelta(minutes=110))=='NEAR'
        assert sla_state(t,t.created_at+timedelta(hours=3))=='BREACHED'
    act(world,tid,'manager','priority',priority='CRITICAL')
    with Session() as db:
        t=db.get(Ticket,tid);assert t.response_deadline-t.created_at==timedelta(minutes=10)
    with Session.begin() as db:
        t=create_ticket(db,db.get(User,world['employee']),TicketCreate(category_id=world['category'],description='Срочное обращение',building='Офис',floor='1',room='1',requested_urgency='URGENT'))
        assert t.priority=='HIGH'

def test_employee_api_isolation_and_internal_notes(client,world):
    tid=make(world)
    with Session.begin() as db:add_comment(db,db.get(User,world['akmal']),tid,Comment(text='TOP INTERNAL',visibility='INTERNAL'))
    login(client,'employee')
    detail=client.get('/api/tickets/'+tid)
    assert detail.status_code==200 and 'TOP INTERNAL' not in detail.text and 'audit' not in detail.json()
    assert client.post('/api/tickets/'+tid+'/comments',json={'text':'secret','visibility':'INTERNAL'}).status_code==403
    assert client.get('/api/users').status_code==403
    assert client.get('/api/analytics').status_code==403
    assert client.get('/api/audit').status_code==403
    assert client.patch('/api/settings',json={'routing_mode':'BROADCAST_ALL'}).status_code==403
    login(client,'employee2')
    assert client.get('/api/tickets/'+tid).status_code==404
    assert client.get('/api/tickets').json()['total']==0

def test_auth_csrf_status_injection(client,world):
    assert client.get('/api/tickets').status_code==401
    login(client,'employee')
    body={'category_id':world['category'],'description':'Не работает компьютер','building':'Офис','floor':'3','room':'305'}
    client.headers.pop('X-CSRF-Token')
    assert client.post('/api/tickets',json=body).status_code==403
    login(client,'employee')
    assert client.post('/api/tickets',json={**body,'status':'CLOSED'}).status_code==422
    assert client.post('/api/tickets',json=body).status_code==201
    assert client.post('/api/auth/login',json={'email':'employee@example.local','password':'bad'}).status_code==401

def test_registration_validation_and_approval(client,world):
    p=Profile(first_name='Алексей',last_name='Иванов',department_id=world['department'],internal_phone='1245',mobile_phone='+998 90 123 45 67')
    with Session.begin() as db:
        u=register_employee(db,9999,9999,'new_user',p);uid=u.id
        assert u.status=='PENDING_APPROVAL' and u.mobile_phone=='+998901234567'
        assert register_employee(db,9999,9999,'changed',p).id==uid
        with pytest.raises(HTTPException):create_ticket(db,u,TicketCreate(category_id=world['category'],description='Test ticket',building='a',floor='1',room='2'))
    login(client)
    assert client.patch('/api/users/'+uid,json={'status':'ACTIVE'}).status_code==200
    with Session() as db:
        assert db.get(User,uid).status=='ACTIVE'
        assert db.scalar(select(Notification).where(Notification.event=='USER_APPROVED'))

def test_filters_and_exports(client,world):
    tid=make(world);login(client,'manager')
    r=client.get('/api/tickets',params={'category_id':world['category'],'q':'компьютер','priority':'NORMAL','status':'NEW','page_size':1})
    assert r.status_code==200 and r.json()['total']==1 and r.json()['items'][0]['id']==tid
    assert client.get('/api/tickets',params={'category_id':world['other_category']}).json()['total']==0
    for fmt in ['csv','xlsx']:
        response=client.get('/api/reports/export',params={'format':fmt})
        assert response.status_code==200
        if fmt=='xlsx':
            from openpyxl import load_workbook
            import io
            wb=load_workbook(io.BytesIO(response.content));assert wb.active.max_row==2
        else:assert 'Не работает компьютер' in response.text
    stats=client.get('/api/analytics').json();assert stats['total']==1 and stats['new']==1

def test_scheduler_autoclose_and_sla_dedup(world):
    from aho.scheduler import tick
    tid=make(world);act(world,tid,'akmal','accept');act(world,tid,'akmal','start');act(world,tid,'akmal','complete',comment='Готово')
    with Session.begin() as db:db.get(Ticket,tid).completed_at=now()-timedelta(hours=80)
    tick()
    with Session() as db:
        t=db.get(Ticket,tid);assert t.status=='CLOSED' and t.closed_by=='SYSTEM' and t.close_reason=='AUTO_CLOSE'
    overdue=make(world)
    with Session.begin() as db:
        t=db.get(Ticket,overdue);t.created_at=now()-timedelta(days=3);t.response_deadline=t.created_at+timedelta(hours=2);t.resolution_deadline=t.created_at+timedelta(hours=24)
    tick();tick()
    with Session() as db:
        events=db.scalars(select(SlaEvent).where(SlaEvent.ticket_id==overdue)).all();assert len(events)==6
        assert db.get(Ticket,overdue).escalated

def test_login_rate_limit(client):
    for _ in range(10):assert client.post('/api/auth/login',json={'email':'nobody','password':'wrong'}).status_code==401
    assert client.post('/api/auth/login',json={'email':'nobody','password':'wrong'}).status_code==429
